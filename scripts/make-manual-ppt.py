#!/usr/bin/env python3
"""
生成「打卡小星星」新家长快速上手手册（pptx）。

用法（需要 python-pptx）：
    python3 scripts/make-manual-ppt.py [输出路径]

为什么把它做成一个脚本而不是手搓一个 pptx：手册里的数字（日薪、球价、
宝可梦总数、等级上限）都会随代码变，脚本至少让"改一处、重新生成一份"成为可能。

配色直接取自 src/app/globals.css 里那套 NES 调色板，手册和产品是一眼能对上的。
字体用微软雅黑：读者多半在 Windows 上用 WPS 或 PowerPoint 打开，这个一定有。
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# ---- 取自 globals.css 的 NES 调色板 ----
SKY = RGBColor(0x5C, 0x94, 0xFC)
RED = RGBColor(0xE4, 0x00, 0x0F)
YELLOW = RGBColor(0xFB, 0xD0, 0x00)
GREEN = RGBColor(0x00, 0xA8, 0x00)
BROWN = RGBColor(0xA8, 0x52, 0x00)
PINK = RGBColor(0xFF, 0x8C, 0xA0)
BLACK = RGBColor(0x10, 0x10, 0x10)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
CREAM = RGBColor(0xFF, 0xF7, 0xE0)

FONT = "微软雅黑"

W, H = Emu(12192000), Emu(6858000)  # 16:9
FOOTER_Y = 17.0  # 页脚上沿，bullets 默认撑到这里为止


def cm(v):
    return Emu(int(v * 360000))


def set_font(run, size, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # 中文要单独指定东亚字体，否则 PowerPoint 会退回宋体
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.makeelement(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1],
            {"typeface": FONT},
        )
        rPr.append(el)


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: [(文字, 字号, 加粗, 颜色, 段前间距pt)]"""
    tb = slide.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (text, size, bold, color, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space)
        p.line_spacing = 1.25
        set_font(p.add_run(), size, bold, color)
        p.runs[0].text = text
    return tb


def box(slide, x, y, w, h, fill, line=BLACK, width=2.5, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, cm(x), cm(y), cm(w), cm(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(width)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def boxtext(shape, lines, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = cm(0.3)
    for i, (text, size, bold, color, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space)
        p.line_spacing = 1.2
        set_font(p.add_run(), size, bold, color)
        p.runs[0].text = text


def arrow(slide, x, y, w, h=0.9, color=BLACK):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, cm(x), cm(y), cm(w), cm(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    return s


def header(slide, no, title, sub=None):
    """每页统一的标题条：左侧一个色块编号 + 标题。"""
    n = box(slide, 1.6, 1.1, 1.5, 1.5, YELLOW)
    boxtext(n, [(no, 26, True, BLACK, 0)])
    textbox(
        slide,
        3.5,
        1.1,
        26,
        1.6,
        [(title, 30, True, BLACK, 0)]
        + ([(sub, 15, False, GREY, 6)] if sub else []),
        anchor=MSO_ANCHOR.MIDDLE,
    )
    ln = box(slide, 1.6, 2.95, 29.5, 0.08, BLACK, BLACK, 0.5)
    ln.line.fill.background()


def footer(slide, text):
    textbox(slide, 1.6, FOOTER_Y, 29.5, 0.8, [(text, 11, False, GREY, 0)])


def bullets(slide, x, y, w, items, size=15, gap=10, h=None):
    """items: [(前缀, 正文)]，前缀会加粗成另一个颜色。

    h 不传就一直撑到页脚上沿。别写死一个大数——放在页面下半部分时
    框子会伸到画布外面去（第一版就是这样，2/3/8 页各超出 4cm）。
    """
    if h is None:
        h = max(1.0, FOOTER_Y - 0.3 - y)
    tb = slide.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = 0
    for i, (lead, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else gap)
        p.line_spacing = 1.35
        if lead:
            set_font(p.add_run(), size, True, RED)
            p.runs[-1].text = lead + "  "
        set_font(p.add_run(), size, False, BLACK)
        p.runs[-1].text = body
    return tb


def build(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ==================== 1. 封面 ====================
    s = blank(prs)
    cover = box(s, 0, 0, W / 360000, H / 360000, SKY, SKY, 0)
    cover.line.fill.background()
    textbox(
        s,
        3.5,
        5.4,
        27,
        8,
        [
            ("打卡小星星", 60, True, WHITE, 0),
            ("一年级孩子的学习打卡 + 奖励系统", 24, False, WHITE, 18),
            ("新家长快速上手（约 10 分钟看完）", 17, False, RGBColor(0xE0, 0xEC, 0xFF), 26),
        ],
    )
    tag = box(s, 3.5, 13.2, 9.2, 1.5, YELLOW)
    boxtext(tag, [("拿到邀请码之后从这里开始", 15, True, BLACK, 0)])

    # ==================== 2. 这是个什么东西 ====================
    s = blank(prs)
    header(s, "0", "它解决的是什么事", "把「今天作业写了吗」从每天的拉锯，变成孩子自己看得见的进度")

    cy = 4.6
    steps = [
        ("完成任务", "读书、练字、\n跳绳……", GREEN),
        ("赚到阳光", "每项任务\n对应多少阳光", YELLOW),
        ("换成奖励", "你和孩子\n一起定的", PINK),
    ]
    x = 2.2
    for i, (t, d, c) in enumerate(steps):
        b = box(s, x, cy, 7.4, 4.2, c)
        boxtext(
            b,
            [
                (t, 21, True, BLACK if c != GREEN else WHITE, 0),
                (d, 13, False, BLACK if c != GREEN else WHITE, 8),
            ],
        )
        if i < 2:
            arrow(s, x + 7.7, cy + 1.75, 1.7, 0.85)
        x += 9.4

    note = box(s, 2.2, 10.0, 29.4, 2.3, CREAM, BROWN, 2)
    boxtext(
        note,
        [
            (
                "剩下的阳光还能投进一条收集线（宝可梦图鉴 / 种花园）——"
                "这条线不花你的钱，但孩子会为了它主动去做任务。",
                15,
                False,
                BLACK,
                0,
            )
        ],
    )

    bullets(
        s,
        2.2,
        13.0,
        29.4,
        [
            ("重要", "奖励清单是你和孩子商量出来的，不是系统给的。系统只负责记账和兑现。"),
            ("不重要", "阳光不是钱。它只在你们家这套约定里有意义，随时可以一起改。"),
        ],
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 3. 注册 ====================
    s = blank(prs)
    header(s, "1", "注册账号", "邀请码一次性，用掉就失效")

    x = 2.2
    for i, (t, d) in enumerate(
        [
            ("打开注册页", "邀请人发给你的链接，\n形如  你的域名/signup"),
            ("填邀请码", "形如 K7QM-3XPT-9WRD\n只能用一次"),
            ("填邮箱 + 密码", "密码至少 8 位。\n这就是你以后的登录账号"),
            ("建孩子档案", "填名字就行。\n后台其它页面才有内容"),
        ]
    ):
        b = box(s, x, 4.6, 6.9, 4.6, WHITE)
        boxtext(
            b,
            [
                (f"{i + 1}", 22, True, RED, 0),
                (t, 17, True, BLACK, 4),
                (d, 12, False, GREY, 6),
            ],
        )
        x += 7.5

    warn = box(s, 2.2, 10.2, 29.4, 2.0, CREAM, RED, 2.5)
    boxtext(
        warn,
        [
            (
                "⚠️  系统跑在邀请人的服务器上，技术上他能看到你家的数据。介意的话先跟他说。",
                16,
                True,
                BLACK,
                0,
            )
        ],
    )

    bullets(
        s,
        2.2,
        12.9,
        29.4,
        [
            ("没有可用的邀请码时", "注册页会自动关闭，直接跳回登录页——不是坏了，是邀请人还没生成码。"),
            ("每个账号相互独立", "你看不到别人家的孩子，别人也看不到你的。"),
        ],
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 4. 选玩法 ====================
    s = blank(prs)
    header(s, "2", "选一个玩法", "注册时就要选，之后也能在「孩子档案」里改")

    a = box(s, 2.2, 4.4, 14.2, 8.4, WHITE, GREEN, 3.5)
    boxtext(a, [("宝可梦图鉴", 26, True, GREEN, 0)])
    rec = box(s, 13.3, 4.0, 2.4, 1.0, GREEN)
    boxtext(rec, [("推荐", 13, True, WHITE, 0)])
    textbox(
        s,
        3.0,
        7.2,
        12.6,
        5.4,
        [
            ("· 用阳光买精灵球，去抓每天遇到的宝可梦", 14, False, BLACK, 0),
            ("· 一共 386 只，按打卡等级一批批解锁", 14, False, BLACK, 8),
            ("· 抓到的永久留在牌库，点进去能看进化路线", 14, False, BLACK, 8),
            ("· 收藏一直往上累积，不会清零", 14, True, BLACK, 8),
            ("内容够玩好几年。孩子看得到长期目标。", 13, False, GREY, 14),
        ],
    )

    b = box(s, 17.4, 4.4, 14.2, 8.4, WHITE, BROWN, 3.5)
    boxtext(b, [("种花园", 26, True, BROWN, 0)])
    textbox(
        s,
        18.2,
        7.2,
        12.6,
        5.4,
        [
            ("· 用阳光买种子，种满一园集齐一套", 14, False, BLACK, 0),
            ("· 收获能拿回阳光，养得越久利息越高", 14, False, BLACK, 8),
            ("· 但每收获一次，花园就清空重来", 14, True, BLACK, 8),
            ("· 植物种类按等级解锁", 14, False, BLACK, 8),
            ("周期短、反馈快，但没有长期收藏感。", 13, False, GREY, 14),
        ],
    )

    textbox(
        s,
        2.2,
        13.6,
        29.4,
        1.6,
        [("两个玩法互斥，一次只能开一个——同时开的话阳光会被稀释，两条线都推不动。", 14, False, GREY, 0)],
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 5. 家长端 / 孩子端（最容易踩坑）====================
    s = blank(prs)
    header(s, "3", "家长端和孩子端，是同一个账号", "这一页是最容易搞混的，值得多看两眼")

    p = box(s, 2.2, 4.4, 13.6, 7.2, WHITE, SKY, 3.5)
    boxtext(p, [("你自己的手机 / 电脑", 20, True, SKY, 0)])
    textbox(
        s,
        3.0,
        6.6,
        12.0,
        4.6,
        [
            ("登录时  ☐ 不勾「这是孩子的设备」", 15, True, BLACK, 0),
            ("→ 进家长后台", 15, False, BLACK, 10),
            ("批任务、管奖励、加临时任务、看统计", 13, False, GREY, 6),
        ],
    )

    k = box(s, 17.9, 4.4, 13.6, 7.2, WHITE, GREEN, 3.5)
    boxtext(k, [("孩子的平板 / 学习机", 20, True, GREEN, 0)])
    textbox(
        s,
        18.7,
        6.6,
        12.0,
        4.6,
        [
            ("登录时  ☑ 勾上「这是孩子的设备」", 15, True, BLACK, 0),
            ("→ 锁在孩子端", 15, False, BLACK, 10),
            ("他点不进后台，改不了自己的分数", 13, False, GREY, 6),
        ],
    )

    esc = box(s, 2.2, 12.4, 29.3, 2.6, CREAM, BROWN, 2.5)
    boxtext(
        esc,
        [
            ("在孩子平板上要进后台怎么办？", 16, True, BROWN, 0),
            ("点孩子端最下面那行很小的「家长登录 →」，重新登录一次（这次别勾）。", 15, False, BLACK, 6),
        ],
        align=PP_ALIGN.LEFT,
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 6. 设任务 ====================
    s = blank(prs)
    header(s, "4", "设孩子的每日任务", "后台 →「任务模板」")

    bullets(
        s,
        2.2,
        4.4,
        16.0,
        [
            ("先做这个", "把示例任务停用，换成孩子真实的每日任务。"),
            ("每项填分值", "比如「读书 20 分钟 = 10 阳光」。"),
            ("三种排期", "「按星期几」自己勾；「上学日」和「休息日」跟着国家放假安排走，含调休补课的周末。"),
            ("每天达标线", "当天挣够这个数算达标；一个月达标够多，月底自动发满勤奖。"),
        ],
        size=15,
    )

    d = box(s, 19.0, 4.4, 12.6, 8.2, CREAM, RED, 3)
    boxtext(
        d,
        [
            ("⚠️ 加任务 = 全场打折", 19, True, RED, 0),
        ],
    )
    textbox(
        s,
        19.8,
        6.4,
        11.0,
        6.0,
        [
            ("所有价格最初都是照着「每天能挣多少」定的。", 13, False, BLACK, 0),
            ("你加一门课，孩子日薪就变高，而礼物和精灵球的价格不会自己跟上——", 13, False, BLACK, 8),
            ("原来要攒 20 天的玩具，可能变成 13 天。", 13, True, BLACK, 8),
            ("后台「经济体检」页会把这件事摊开，并给一个带预览的「一键校准」。", 13, False, BLACK, 10),
            ("改不改由你决定。", 13, True, GREEN, 8),
        ],
    )

    textbox(
        s,
        2.2,
        13.4,
        16.0,
        2.0,
        [("系统自动发的那些（满勤奖、图鉴奖励、刷新费）会自己跟着日薪走，不用管。", 13, False, GREY, 0)],
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 7. 设礼物 ====================
    s = blank(prs)
    header(s, "5", "设奖励清单", "后台 →「礼物」。这一步建议和孩子一起做")

    textbox(
        s,
        2.2,
        3.6,
        29.4,
        1.0,
        [
            (
                "价格别照抄别人家的——系统按「几天工资」判断贵不贵，"
                "而每家孩子的日薪不一样。下面是四档参考区间：",
                14,
                False,
                GREY,
                0,
            )
        ],
    )

    # 四档和冷却天数的对应关系直接取自 lib/economy.ts 的 PRICE_BANDS
    rows = [
        ("日常小奖励", "冷却 ≤ 1 天", "0.4 ~ 1.2 天工资", "看半小时动画片、今晚吃什么我说了算", GREEN),
        ("周奖励", "冷却 2 ~ 7 天", "1.2 ~ 7 天工资", "买一本新书、周末去趟游乐场", YELLOW),
        ("半月奖励", "冷却 8 ~ 14 天", "3 ~ 8 天工资", "一个小玩具", PINK),
        ("攒很久的大奖励", "冷却 ≥ 15 天", "14 ~ 28 天工资", "一套乐高、去一次迪士尼", BROWN),
    ]
    y = 5.0
    for t, cd, price, eg, c in rows:
        bar = box(s, 2.2, y, 1.0, 2.1, c)
        bar.line.fill.background()
        box(s, 3.2, y, 28.4, 2.1, WHITE)
        textbox(
            s,
            4.0,
            y + 0.25,
            26.8,
            1.6,
            [
                (f"{t}　·　{cd}　·　{price}", 15, True, BLACK, 0),
                (f"例如：{eg}", 12, False, GREY, 3),
            ],
        )
        y += 2.35

    tip = box(s, 2.2, 14.6, 29.4, 2.6, CREAM, RED, 2.5)
    boxtext(
        tip,
        [
            ("档位是从冷却天数推出来的，所以贵礼物一定要设冷却", 16, True, RED, 0),
            (
                "不设冷却的话，系统没法判断这是「日常小奖励」还是「你忘了设冷却的大奖」，"
                "体检时会直接跳过它；孩子那边也可能把阳光全砸在同一件事上。",
                13,
                False,
                BLACK,
                6,
            ),
        ],
        align=PP_ALIGN.LEFT,
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 8. 孩子怎么用 ====================
    s = blank(prs)
    header(s, "6", "把孩子端交给孩子", "后台 →「孩子档案」，那里有一条完整链接")

    x = 2.2
    for i, (t, d) in enumerate(
        [
            ("复制链接", "「孩子档案」里每个孩子\n下面就是完整地址，\n点一下能直接打开"),
            ("在孩子设备上打开", "登录时勾上\n「这是孩子的设备」"),
            ("加到主屏幕", "用浏览器的\n「添加到主屏幕」，\n之后点图标就进"),
        ]
    ):
        b = box(s, x, 4.5, 9.5, 4.4, WHITE)
        boxtext(
            b,
            [(f"{i + 1}", 20, True, RED, 0), (t, 17, True, BLACK, 4), (d, 12, False, GREY, 6)],
        )
        x += 10.1

    dis = box(s, 2.2, 9.8, 29.4, 3.0, CREAM, RED, 2.5)
    boxtext(
        dis,
        [
            ("孩子需要知道的一条规则", 17, True, RED, 0),
            (
                "任务没全部完成的日子，会跑掉一只已经抓到的宝可梦（花园主题则是被吃掉一棵植物）。"
                "当天只要抓到过宝可梦就免疫。不开 App 也躲不掉——第二天打开会一次性补算。",
                14,
                False,
                BLACK,
                6,
            ),
        ],
        align=PP_ALIGN.LEFT,
    )

    bullets(
        s,
        2.2,
        13.4,
        29.4,
        [
            ("孩子端每页一屏装得下", "平板和手机都适配过，不用滚动，字号是按认字阶段调的，全部带拼音。"),
            ("家长也能用手机替他打卡", "手机上打开孩子端，任务列表占满整屏。"),
        ],
        size=14,
    )
    footer(s, "打卡小星星 · 新家长快速上手")

    # ==================== 9. 常见问题 ====================
    s = blank(prs)
    header(s, "7", "几个大概率会问到的问题")

    qa = [
        ("批错了怎么办？", "后台「打卡记录」里可以撤销，阳光会扣回去，流水上留一条记录。"),
        ("孩子今天额外做了事？", "后台加一条「临时任务」，比直接手动加分好——手动加分不算进等级。"),
        ("等级有什么用？", "共 15 级，按累计挣到的阳光升。等级解锁宝可梦的数量和地区、花园的植物种类。"),
        ("阳光会不会通货膨胀？", "会。所以有「经济体检」——它按当前日薪算出每样东西值几天工资，越界的标出来。"),
        ("孩子能自己改分数吗？", "不能。勾了「这是孩子的设备」的会话进不了后台。"),
        ("能几个孩子一起用吗？", "可以，后台顶栏切换。数据各自独立。"),
    ]
    y = 4.3
    for i, (q, a) in enumerate(qa):
        c = LIGHT if i % 2 == 0 else WHITE
        bg = box(s, 2.2, y, 29.4, 1.95, c, c, 0.5)
        bg.line.fill.background()
        textbox(
            s,
            2.9,
            y + 0.28,
            28.0,
            1.5,
            [(q, 15, True, BLACK, 0), (a, 13, False, GREY, 3)],
        )
        y += 2.05

    footer(s, "打卡小星星 · 新家长快速上手　|　有问题直接问邀请你的那位家长")

    prs.save(path)
    return len(prs.slides._sldIdLst)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "打卡小星星-新家长快速上手.pptx"
    n = build(out)
    print(f"已生成 {out}（{n} 页）")
